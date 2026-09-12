# from docx import Document

# from backend.app.ingestion.loaders.docx_loader import load_docx
# from backend.app.ingestion.models.source_document import SourceDocument
# from backend.app.ingestion.models.content import TableContent

# def test_load_docx_with_paragraphs_and_tables(tmp_path):
#     file_path = tmp_path / "Sonal_Anand_Resume_Skwad.docx"

#     document = Document()

#     document.add_paragraph("FULL STACK DEVELOPER ")
#     document.add_paragraph("React Native mobile app")

#     table = document.add_table(rows=1, cols=3)

#     header = table.rows[0].cells
#     header[0].text = "Component"
#     header[1].text = "Status"
#     header[2].text = "Owner"

#     row = table.add_row().cells
#     row[0].text = "Backend"
#     row[1].text = "Done"
#     row[2].text = "Sonal"

#     document.save(file_path)

#     documents = load_docx(
#         str(file_path),
#         document_id="docx-1",
#         filename="Sonal_Anand_Resume_Skwad.docx",
#     )

#     assert len(documents) == 2

#     # Paragraph document
#     text_document = documents[0]

#     assert isinstance(text_document.content, TableContent)
#     assert text_document.source_type == "docx"
#     assert text_document.metadata["content_type"] == "text"
#     assert ("FULL STACK DEVELOPER" in text_document.content.text)

#     # Table document
#     table_document = documents[1]

#     assert isinstance(table_document.content, TableContent)
#     assert table_document.source_type == "docx"
#     assert table_document.metadata["content_type"] == "table"
#     assert table_document.metadata["table_index"] == 1

#     assert ("Component | Status | Owner" in table_document.content.text)
#     assert ("Backend | Done | Sonal" in table_document.content.text)

from pathlib import Path
from PIL import Image
import pytest
from docx import Document
from docx.shared import Inches

from backend.app.ingestion.loaders.docx_loader import (
    load_docx,
)
from backend.app.ingestion.models.content import (
    ImageContent,
    TableContent,
    TextContent,
)


PNG_BYTES = bytes(
    [
        137,
        80,
        78,
        71,
        13,
        10,
        26,
        10,
        0,
        0,
        0,
        13,
        73,
        72,
        68,
        82,
        0,
        0,
        0,
        1,
        0,
        0,
        0,
        1,
        8,
        4,
        0,
        0,
        0,
        181,
        28,
        12,
        2,
        0,
        0,
        0,
        11,
        73,
        68,
        65,
        84,
        120,
        156,
        99,
        248,
        207,
        192,
        240,
        31,
        0,
        3,
        3,
        1,
        255,
        167,
        84,
        53,
        0,
        0,
        0,
        0,
        73,
        69,
        78,
        68,
        174,
        66,
        96,
        130,
    ]
)


def create_test_docx(
    docx_path: Path,
) -> None:
    image_path = (
        docx_path.parent / "test-image.png"
    )

    image = Image.new(
        "RGB",
        (100, 100),
        "white",
    )

    image.save(
        image_path,
        format="PNG",
    )

    document = Document()

    document.add_heading(
        "Knowledge Intelligence Assistant",
        level=1,
    )

    paragraph = document.add_paragraph(
        "This document contains an embedded image "
        "and a data table."
    )

    run = paragraph.add_run()

    run.add_picture(
        str(image_path),
        width=Inches(1),
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Technology"
    table.cell(1, 0).text = "Backend"
    table.cell(1, 1).text = "FastAPI"

    document.save(
        str(docx_path)
    )

def test_load_docx_extracts_content(
    tmp_path,
):
    docx_path = tmp_path / "sample.docx"
    output_dir = tmp_path / "output"

    create_test_docx(
        docx_path
    )

    result = load_docx(
        file_path=str(docx_path),
        document_id="docx-123",
        output_dir=str(output_dir),
    )

    element_types = {
        element.element_type
        for element in result.elements
    }

    assert "document" in element_types
    assert "paragraph" in element_types
    assert "table" in element_types
    assert "image" in element_types


def test_load_docx_extracts_embedded_image(
    tmp_path,
):
    docx_path = tmp_path / "sample.docx"
    output_dir = tmp_path / "output"

    create_test_docx(
        docx_path
    )

    result = load_docx(
        file_path=str(docx_path),
        document_id="docx-123",
        output_dir=str(output_dir),
    )

    images = [
        element
        for element in result.elements
        if element.element_type == "image"
    ]

    assert len(images) == 1

    image = images[0]

    assert isinstance(
        image.content,
        ImageContent,
    )

    assert Path(
        image.content.path
    ).exists()


def test_load_docx_extracts_table(
    tmp_path,
):
    docx_path = tmp_path / "sample.docx"
    output_dir = tmp_path / "output"

    create_test_docx(
        docx_path
    )

    result = load_docx(
        file_path=str(docx_path),
        document_id="docx-123",
        output_dir=str(output_dir),
    )

    tables = [
        element
        for element in result.elements
        if element.element_type == "table"
    ]

    assert len(tables) == 1

    table = tables[0]

    assert isinstance(
        table.content,
        TableContent,
    )

    assert "Component" in table.content.text
    assert "FastAPI" in table.content.text


def test_load_docx_creates_image_relationship(
    tmp_path,
):
    docx_path = tmp_path / "sample.docx"
    output_dir = tmp_path / "output"

    create_test_docx(
        docx_path
    )

    result = load_docx(
        file_path=str(docx_path),
        document_id="docx-123",
        output_dir=str(output_dir),
    )

    image_relationships = [
        relationship
        for relationship in result.relationships
        if (
            relationship.relationship_type
            == "contains"
            and relationship.metadata.get(
                "relationship"
            )
            == "embedded_image"
        )
    ]

    assert image_relationships

    image_ids = {
        element.element_id
        for element in result.elements
        if element.element_type == "image"
    }

    assert (
        image_relationships[0].target_element_id
        in image_ids
    )


def test_load_docx_creates_document_relationships(
    tmp_path,
):
    docx_path = tmp_path / "sample.docx"
    output_dir = tmp_path / "output"

    create_test_docx(
        docx_path
    )

    result = load_docx(
        file_path=str(docx_path),
        document_id="docx-123",
        output_dir=str(output_dir),
    )

    relationships = [
        relationship
        for relationship in result.relationships
        if relationship.source_element_id
        == "document"
    ]

    assert relationships

    element_ids = {
        element.element_id
        for element in result.elements
    }

    for relationship in relationships:
        assert (
            relationship.target_element_id
            in element_ids
        )


def test_load_docx_missing_file(
    tmp_path,
):
    with pytest.raises(
        FileNotFoundError
    ):
        load_docx(
            file_path=str(
                tmp_path / "missing.docx"
            ),
            document_id="docx-123",
            output_dir=str(
                tmp_path / "output"
            ),
        )


def test_load_docx_rejects_non_docx(
    tmp_path,
):
    file_path = tmp_path / "sample.pdf"

    file_path.write_bytes(
        b"fake"
    )

    with pytest.raises(
        ValueError
    ):
        load_docx(
            file_path=str(file_path),
            document_id="docx-123",
            output_dir=str(
                tmp_path / "output"
            ),
        )