import base64
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from backend.app.ingestion.loaders.pptx_loader import (
    load_pptx,
)
from backend.app.ingestion.models.content import (
    ImageContent,
    TableContent,
    TextContent,
)


# Tiny valid 1x1 PNG.
PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAAB"
    "CAQAAAC1HAwCAAAAC0lEQVR42mNk+A8"
    "AAQUBAScY42YAAAAASUVORK5CYII="
)


def create_test_pptx(
    pptx_path: Path,
) -> None:
    presentation = Presentation()

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    # Text
    text_box = slide.shapes.add_textbox(
        Inches(1),
        Inches(1),
        Inches(5),
        Inches(1),
    )

    text_box.text = "Knowledge Intelligence Assistant"

    # Image
    image_path = pptx_path.parent / "test-image.png"
    image_path.write_bytes(PNG_BYTES)

    slide.shapes.add_picture(
        str(image_path),
        Inches(1),
        Inches(2),
        width=Inches(2),
        height=Inches(2),
    )

    # Table
    table_shape = slide.shapes.add_table(
        rows=2,
        cols=2,
        left=Inches(4),
        top=Inches(2),
        width=Inches(3),
        height=Inches(2),
    )

    table = table_shape.table

    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Technology"
    table.cell(1, 0).text = "Backend"
    table.cell(1, 1).text = "FastAPI"

    presentation.save(pptx_path)


def test_load_pptx_extracts_text_image_and_table(
    tmp_path,
):
    pptx_path = tmp_path / "sample.pptx"
    output_dir = tmp_path / "output"

    create_test_pptx(pptx_path)

    result = load_pptx(
        file_path=str(pptx_path),
        document_id="pptx-123",
        output_dir=str(output_dir),
    )

    assert result.document_id == "pptx-123"

    element_types = [
        element.element_type
        for element in result.elements
    ]

    assert "slide" in element_types
    assert "text" in element_types
    assert "image" in element_types
    assert "table" in element_types

    text_elements = [
        element
        for element in result.elements
        if element.element_type == "text"
    ]

    assert any(
        isinstance(element.content, TextContent)
        and (
            element.content.text
            == "Knowledge Intelligence Assistant"
        )
        for element in text_elements
    )

    image_elements = [
        element
        for element in result.elements
        if element.element_type == "image"
    ]

    assert len(image_elements) == 1
    assert isinstance(
        image_elements[0].content,
        ImageContent,
    )
    assert Path(
        image_elements[0].content.path
    ).exists()

    table_elements = [
        element
        for element in result.elements
        if element.element_type == "table"
    ]

    assert len(table_elements) == 1
    assert isinstance(
        table_elements[0].content,
        TableContent,
    )

    assert "Component" in table_elements[0].content.text
    assert "FastAPI" in table_elements[0].content.text


def test_load_pptx_creates_slide_relationships(
    tmp_path,
):
    pptx_path = tmp_path / "sample.pptx"
    output_dir = tmp_path / "output"

    create_test_pptx(pptx_path)

    result = load_pptx(
        file_path=str(pptx_path),
        document_id="pptx-123",
        output_dir=str(output_dir),
    )

    contains_relationships = [
        relationship
        for relationship in result.relationships
        if relationship.relationship_type
        == "contains"
    ]

    assert contains_relationships

    target_ids = {
        relationship.target_element_id
        for relationship in contains_relationships
    }

    assert "slide-1" not in target_ids
    assert any(
        target_id.startswith("slide-1-image-")
        for target_id in target_ids
    )
    assert any(
        target_id.startswith("slide-1-table-")
        for target_id in target_ids
    )
    assert any(
        target_id.startswith("slide-1-text-")
        for target_id in target_ids
    )


def test_load_pptx_preserves_shape_geometry(
    tmp_path,
):
    pptx_path = tmp_path / "sample.pptx"
    output_dir = tmp_path / "output"

    create_test_pptx(pptx_path)

    result = load_pptx(
        file_path=str(pptx_path),
        document_id="pptx-123",
        output_dir=str(output_dir),
    )

    image_elements = [
        element
        for element in result.elements
        if element.element_type == "image"
    ]

    assert len(image_elements) == 1

    bbox = image_elements[0].metadata["bbox"]

    assert bbox["left"] > 0
    assert bbox["top"] > 0
    assert bbox["width"] > 0
    assert bbox["height"] > 0


def test_load_pptx_missing_file(tmp_path):
    missing_path = tmp_path / "missing.pptx"

    with pytest.raises(FileNotFoundError):
        load_pptx(
            file_path=str(missing_path),
            document_id="pptx-123",
            output_dir=str(tmp_path / "output"),
        )


def test_load_pptx_rejects_unsupported_format(
    tmp_path,
):
    file_path = tmp_path / "sample.pdf"
    file_path.write_bytes(b"fake")

    with pytest.raises(ValueError):
        load_pptx(
            file_path=str(file_path),
            document_id="pptx-123",
            output_dir=str(tmp_path / "output"),
        )