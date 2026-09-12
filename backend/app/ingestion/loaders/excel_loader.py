# import pandas as pd

# from backend.app.ingestion.models.source_document import SourceDocument
# from backend.app.ingestion.models.content import (
#     TableContent,
#     TextContent,
# )


# def load_excel(
#     file_path: str,
#     document_id: str,
#     file_name: str,
# ) -> list[SourceDocument]:
#     workbook = pd.ExcelFile(file_path)

#     documents = []

#     for sheet_name in workbook.sheet_names:
#         dataframe = pd.read_excel(
#             workbook,
#             sheet_name=sheet_name,
#         )

#         if dataframe.empty:
#             continue

#         rows = []

#         for _, row in dataframe.iterrows():
#             row_content = " | ".join(
#                 f"{column}: {value}"
#                 for column, value in row.items()
#                 if pd.notna(value)
#             )

#             if row_content.strip():
#                 rows.append(row_content)

#         if not rows:
#             continue

#         content = "\n".join(rows)

#         documents.append(
#             SourceDocument(
#                 content=TableContent(text=content),
#                 document_id=document_id,
#                 source_type="excel",
#                 metadata={
#                     "file_name": file_name,
#                     "content_type": "table",
#                     "sheet_name": sheet_name,
#                     "row_count": len(dataframe),
#                     "column_count": len(dataframe.columns),
#                 },
#             )
#         )

#     return documents

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.app.ingestion.models.content import (
    ImageContent,
    TableContent,
)
from backend.app.ingestion.models.element_relationship import (
    ElementRelationship,
)
from backend.app.ingestion.models.ingestion_result import (
    IngestionResult,
)
from backend.app.ingestion.models.source_element import (
    SourceElement,
)


SUPPORTED_EXCEL_EXTENSIONS = {
    ".xlsx",
    ".xlsm",
}


def _cell_range_from_chart_reference(
    reference: Any,
) -> str | None:
    """
    Extract the worksheet/cell range formula from an
    openpyxl chart reference object.
    """
    if reference is None:
        return None

    num_ref = getattr(
        reference,
        "numRef",
        None,
    )

    if num_ref is None:
        return None

    formula = getattr(
        num_ref,
        "f",
        None,
    )

    if not formula:
        return None

    return str(formula)


def _get_chart_references(
    chart: Any,
) -> list[str]:
    references: list[str] = []

    for series in getattr(
        chart,
        "series",
        [],
    ):
        # Numeric/value series.
        value_ref = _cell_range_from_chart_reference(
            getattr(
                series,
                "val",
                None,
            )
        )

        if value_ref:
            references.append(
                value_ref
            )

        # Category references.
        category = getattr(
            series,
            "cat",
            None,
        )

        category_ref = _cell_range_from_chart_reference(
            category
        )

        if category_ref:
            references.append(
                category_ref
            )

        # XY charts may store x/y references separately.
        x_value_ref = _cell_range_from_chart_reference(
            getattr(
                series,
                "xVal",
                None,
            )
        )

        if x_value_ref:
            references.append(
                x_value_ref
            )

        y_value_ref = _cell_range_from_chart_reference(
            getattr(
                series,
                "yVal",
                None,
            )
        )

        if y_value_ref:
            references.append(
                y_value_ref
            )

    return list(
        dict.fromkeys(references)
    )


def _get_image_anchor(
    image: Any,
) -> dict[str, Any]:
    anchor = getattr(
        image,
        "anchor",
        None,
    )

    if anchor is None:
        return {}

    marker = getattr(
        anchor,
        "_from",
        None,
    )

    if marker is None:
        return {}

    return {
        "row": getattr(
            marker,
            "row",
            None,
        ),
        "column": getattr(
            marker,
            "col",
            None,
        ),
    }

def _get_chart_title(
    chart: Any,
) -> str | None:
    title = getattr(
        chart,
        "title",
        None,
    )

    if title is None:
        return None

    text = getattr(
        getattr(
            getattr(
                title,
                "tx",
                None,
            ),
            "rich",
            None,
        ),
        "p",
        None,
    )

    if not text:
        return None

    parts: list[str] = []

    for paragraph in text:
        for run in getattr(
            paragraph,
            "r",
            [],
        ):
            run_text = getattr(
                run,
                "t",
                None,
            )

            if run_text:
                parts.append(
                    run_text
                )

    value = "".join(parts).strip()

    return value or None

def load_excel(
    file_path: str,
    document_id: str,
    output_dir: str,
) -> IngestionResult:
    excel_path = Path(
        file_path
    )

    if not excel_path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {file_path}"
        )

    if not excel_path.is_file():
        raise ValueError(
            f"Excel path is not a file: {file_path}"
        )

    extension = (
        excel_path.suffix.lower()
    )

    if extension not in SUPPORTED_EXCEL_EXTENSIONS:
        raise ValueError(
            f"Unsupported Excel format: {extension}"
        )

    output_path = Path(
        output_dir
    )

    output_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = load_workbook(
        filename=str(excel_path),
        data_only=False,
    )

    elements: list[SourceElement] = []
    relationships: list[
        ElementRelationship
    ] = []

    workbook_element = SourceElement(
        element_id="workbook",
        document_id=document_id,
        element_type="workbook",
        content=TableContent(
            text=""
        ),
        metadata={
            "file_name": excel_path.name,
            "content_type": "workbook",
        },
    )

    elements.append(
        workbook_element
    )

    for sheet_index, worksheet in enumerate(
        workbook.worksheets,
        start=1,
    ):
        sheet_element_id = (
            f"sheet-{sheet_index}"
        )

        sheet_element = SourceElement(
            element_id=sheet_element_id,
            document_id=document_id,
            element_type="sheet",
            content=TableContent(
                text=""
            ),
            metadata={
                "file_name": excel_path.name,
                "content_type": "sheet",
                "sheet_name": worksheet.title,
                "sheet_index": sheet_index,
            },
        )

        elements.append(
            sheet_element
        )

        relationships.append(
            ElementRelationship(
                source_element_id="workbook",
                relationship_type="contains",
                target_element_id=sheet_element_id,
                metadata={
                    "sheet_name": worksheet.title,
                },
            )
        )

        # -----------------------------------------------------
        # Worksheet data
        # -----------------------------------------------------
        rows = []

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = [
                "" if value is None else str(value)
                for value in row
            ]

            if any(
                value.strip()
                for value in values
            ):
                rows.append(values)

        if rows:
            table_text = "\n".join(
                " | ".join(row)
                for row in rows
            )

            table_element_id = (
                f"sheet-{sheet_index}-table"
            )

            table_element = SourceElement(
                element_id=table_element_id,
                document_id=document_id,
                element_type="table",
                content=TableContent(
                    text=table_text
                ),
                metadata={
                    "file_name": excel_path.name,
                    "content_type": "table",
                    "sheet_name": worksheet.title,
                    "sheet_index": sheet_index,
                    "row_count": len(rows),
                    "column_count": max(
                        (
                            len(row)
                            for row in rows
                        ),
                        default=0,
                    ),
                    "cell_range": worksheet.calculate_dimension(),
                },
            )

            elements.append(
                table_element
            )

            relationships.append(
                ElementRelationship(
                    source_element_id=sheet_element_id,
                    relationship_type="contains",
                    target_element_id=table_element_id,
                    metadata={
                        "sheet_name": worksheet.title,
                    },
                )
            )

        # -----------------------------------------------------
        # Embedded images
        # -----------------------------------------------------
        for image_index, image in enumerate(
            getattr(
                worksheet,
                "_images",
                [],
            ),
            start=1,
        ):
            image_extension = (
                getattr(
                    image,
                    "format",
                    None,
                )
                or "png"
            ).lower()

            image_path = (
                output_path
                / (
                    f"{document_id}_"
                    f"sheet_{sheet_index}_"
                    f"image_{image_index}."
                    f"{image_extension}"
                )
            )

            image_data = None

            ref = getattr(
                image,
                "_data",
                None,
            )

            if callable(ref):
                image_data = ref()

            if not image_data:
                continue

            image_path.write_bytes(
                image_data
            )

            element_id = (
                f"sheet-{sheet_index}-"
                f"image-{image_index}"
            )

            anchor = _get_image_anchor(
                image
            )

            image_element = SourceElement(
                element_id=element_id,
                document_id=document_id,
                element_type="image",
                content=ImageContent(
                    path=str(image_path)
                ),
                metadata={
                    "file_name": excel_path.name,
                    "content_type": "image",
                    "sheet_name": worksheet.title,
                    "sheet_index": sheet_index,
                    "image_index": image_index,
                    "image_path": str(image_path),
                    "image_extension": image_extension,
                    "anchor": anchor,
                },
            )

            elements.append(
                image_element
            )

            relationships.append(
                ElementRelationship(
                    source_element_id=sheet_element_id,
                    relationship_type="contains",
                    target_element_id=element_id,
                    metadata={
                        "sheet_name": worksheet.title,
                    },
                )
            )

        # -----------------------------------------------------
        # Charts
        # -----------------------------------------------------
        for chart_index, chart in enumerate(
            getattr(
                worksheet,
                "_charts",
                [],
            ),
            start=1,
        ):
            chart_element_id = (
                f"sheet-{sheet_index}-"
                f"chart-{chart_index}"
            )

            chart_title = _get_chart_title(
                chart
            )

            source_references = (
                _get_chart_references(
                    chart
                )
            )

            anchor = getattr(
                chart,
                "anchor",
                None,
            )

            chart_metadata = {
                "file_name": excel_path.name,
                "content_type": "chart",
                "sheet_name": worksheet.title,
                "sheet_index": sheet_index,
                "chart_index": chart_index,
                "chart_type": type(
                    chart
                ).__name__,
                "chart_title": chart_title,
                "source_references": source_references,
            }

            if anchor is not None:
                chart_metadata[
                    "anchor"
                ] = {
                    "from_row": getattr(
                        getattr(
                            anchor,
                            "_from",
                            None,
                        ),
                        "row",
                        None,
                    ),
                    "from_column": getattr(
                        getattr(
                            anchor,
                            "_from",
                            None,
                        ),
                        "col",
                        None,
                    ),
                }

            chart_element = SourceElement(
                element_id=chart_element_id,
                document_id=document_id,
                element_type="chart",
                content=TableContent(
                    text=chart_title or ""
                ),
                metadata=chart_metadata,
            )

            elements.append(
                chart_element
            )

            relationships.append(
                ElementRelationship(
                    source_element_id=sheet_element_id,
                    relationship_type="contains",
                    target_element_id=chart_element_id,
                    metadata={
                        "sheet_name": worksheet.title,
                    },
                )
            )

            # Preserve chart → source-data lineage.
            for source_reference in source_references:
                relationships.append(
                    ElementRelationship(
                        source_element_id=chart_element_id,
                        relationship_type="derived_from",
                        target_element_id=(
                            table_element_id
                            if rows
                            else sheet_element_id
                        ),
                        metadata={
                            "source_reference": (
                                source_reference
                            ),
                        },
                    )
                )

    workbook.close()

    return IngestionResult(
        document_id=document_id,
        elements=elements,
        relationships=relationships,
    )