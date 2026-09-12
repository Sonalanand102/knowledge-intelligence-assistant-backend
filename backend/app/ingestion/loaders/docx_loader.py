# from docx import Document

# from backend.app.ingestion.models.source_document import SourceDocument
# from backend.app.ingestion.models.content import TableContent


# def load_docx(
#     file_path: str,
#     document_id: str,
#     filename: str,
# ) -> list[SourceDocument]:
#     document = Document(file_path)

#     documents = []

#     # Extract paragraphs
#     content_parts = []

#     for paragraph in document.paragraphs:
#         if paragraph.text.strip():
#             content_parts.append(paragraph.text.strip())

#     content = "\n\n".join(content_parts)
    
#     if content_parts:
#         documents.append(
#             SourceDocument(
#                 content=TableContent(text=content),
#                 document_id=document_id,
#                 source_type="docx",
#                 metadata={
#                     "filename": filename,
#                     "content_type": "text",
#                 },
#             )
#         )

#     # Extract tables
#     for table_index, table in enumerate(document.tables, start=1):
#         rows = []

#         for row in table.rows:
#             cells = [cell.text.strip() for cell in row.cells]
#             rows.append(" | ".join(cells))

#         table_content = "\n".join(rows)

#         if table_content.strip():
#             documents.append(
#                 SourceDocument(
#                     content=TableContent(text=table_content),
#                     document_id=document_id,
#                     source_type="docx",
#                     metadata={
#                         "filename": filename,
#                         "content_type": "table",
#                         "table_index": table_index,
#                     },
#                 )
#             )

#     return documents

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document

from backend.app.ingestion.models.content import (
    ImageContent,
    TableContent,
    TextContent,
)
from backend.app.ingestion.models.element_relationship import (
    ElementRelationship,
)
from backend.app.ingestion.models.ingestion_result import (
    IngestionResult,
)
from backend.app.ingestion.models.source_element import (
    SourceElement,
)


WORD_NAMESPACE = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

RELATIONSHIP_NAMESPACE = {
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}

IMAGE_RELATIONSHIP_TYPE = (
    "http://schemas.openxmlformats.org/"
    "officeDocument/2006/relationships/image"
)


def _extract_hyperlinks(
    paragraph,
) -> list[tuple[str, str]]:
    hyperlinks = []

    for hyperlink in paragraph._p.xpath(".//w:hyperlink"):
        relationship_id = hyperlink.get(
            f"{{{WORD_NAMESPACE['r']}}}id"
        )

        text = "".join(
            node.text or ""
            for node in hyperlink.xpath(".//w:t")
        ).strip()

        if relationship_id and text:
            hyperlinks.append(
                (relationship_id, text)
            )

    return hyperlinks


def _get_relationship_targets(
    docx_path: Path,
) -> dict[str, str]:
    targets: dict[str, str] = {}

    with ZipFile(docx_path) as archive:
        relationships_path = (
            "word/_rels/document.xml.rels"
        )

        if relationships_path not in archive.namelist():
            return targets

        root = ET.fromstring(
            archive.read(
                relationships_path
            )
        )

        for relationship in root:
            relationship_id = relationship.get(
                "Id"
            )
            target = relationship.get(
                "Target"
            )

            if relationship_id and target:
                targets[relationship_id] = target

    return targets


def _extract_embedded_images(
    docx_path: Path,
    output_dir: Path,
    document_id: str,
) -> dict[str, SourceElement]:
    image_elements: dict[str, SourceElement] = {}

    relationship_targets = _get_relationship_targets(
        docx_path
    )

    with ZipFile(docx_path) as archive:
        media_files = [
            name
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]

        for index, media_name in enumerate(
            media_files,
            start=1,
        ):
            image_data = archive.read(
                media_name
            )

            extension = Path(
                media_name
            ).suffix.lower().lstrip(
                "."
            ) or "png"

            image_path = (
                output_dir
                / (
                    f"{document_id}_"
                    f"image_{index}."
                    f"{extension}"
                )
            )

            image_path.write_bytes(
                image_data
            )

            element_id = (
                f"image-{index}"
            )

            image_element = SourceElement(
                element_id=element_id,
                document_id=document_id,
                element_type="image",
                content=ImageContent(
                    path=str(image_path)
                ),
                metadata={
                    "content_type": "image",
                    "file_path": str(image_path),
                    "extension": extension,
                    "source_member": media_name,
                },
            )

            image_elements[
                media_name
            ] = image_element

    # Keep relationship_targets available for future
    # precise paragraph → image mapping.
    _ = relationship_targets

    return image_elements


def _extract_paragraph_image_relationship_ids(
    paragraph,
) -> list[str]:
    relationship_ids = []

    for drawing in paragraph._p.xpath(
        ".//a:blip"
    ):
        relationship_id = drawing.get(
            "{http://schemas.openxmlformats.org/"
            "officeDocument/2006/relationships}embed"
        )

        if relationship_id:
            relationship_ids.append(
                relationship_id
            )

    return relationship_ids


def load_docx(
    file_path: str,
    document_id: str,
    output_dir: str,
) -> IngestionResult:
    docx_path = Path(file_path)

    if not docx_path.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {file_path}"
        )

    if not docx_path.is_file():
        raise ValueError(
            f"DOCX path is not a file: {file_path}"
        )

    if docx_path.suffix.lower() != ".docx":
        raise ValueError(
            f"Unsupported document format: "
            f"{docx_path.suffix}"
        )

    output_path = Path(output_dir)

    output_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    document = Document(
        str(docx_path)
    )

    elements: list[SourceElement] = []
    relationships: list[
        ElementRelationship
    ] = []

    image_elements = _extract_embedded_images(
        docx_path=docx_path,
        output_dir=output_path,
        document_id=document_id,
    )

    elements.extend(
        image_elements.values()
    )

    # ---------------------------------------------------------
    # Paragraphs
    # ---------------------------------------------------------
    paragraph_index = 0

    relationship_targets = _get_relationship_targets(
        docx_path
    )

    media_by_target = {}

    for media_name, image_element in image_elements.items():
        relative_target = media_name.replace(
            "word/",
            "",
            1,
        )

        media_by_target[
            relative_target
        ] = image_element

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        paragraph_index += 1

        element_id = (
            f"paragraph-{paragraph_index}"
        )

        paragraph_element = SourceElement(
            element_id=element_id,
            document_id=document_id,
            element_type="paragraph",
            content=TextContent(
                text=text
            ),
            metadata={
                "file_name": docx_path.name,
                "content_type": "text",
                "element_index": paragraph_index,
                "style": paragraph.style.name,
            },
        )

        elements.append(
            paragraph_element
        )

        hyperlink_data = _extract_hyperlinks(
            paragraph
        )

        for link_index, (
            relationship_id,
            link_text,
        ) in enumerate(
            hyperlink_data,
            start=1,
        ):
            target = relationship_targets.get(
                relationship_id
            )

            if not target:
                continue

            link_element_id = (
                f"paragraph-{paragraph_index}-"
                f"link-{link_index}"
            )

            link_element = SourceElement(
                element_id=link_element_id,
                document_id=document_id,
                element_type="link",
                content=TextContent(
                    text=link_text
                ),
                metadata={
                    "file_name": docx_path.name,
                    "content_type": "link",
                    "url": target,
                    "paragraph_index": paragraph_index,
                },
            )

            elements.append(
                link_element
            )

            relationships.append(
                ElementRelationship(
                    source_element_id=(
                        element_id
                    ),
                    relationship_type="contains",
                    target_element_id=(
                        link_element_id
                    ),
                )
            )

        # Paragraph → image relationships.
        image_relationship_ids = (
            _extract_paragraph_image_relationship_ids(
                paragraph
            )
        )

        for relationship_id in (
            image_relationship_ids
        ):
            target = relationship_targets.get(
                relationship_id
            )

            if not target:
                continue

            normalized_target = target.replace(
                "\\",
                "/",
            )

            target_image = media_by_target.get(
                normalized_target
            )

            if not target_image:
                continue

            relationships.append(
                ElementRelationship(
                    source_element_id=(
                        element_id
                    ),
                    relationship_type="contains",
                    target_element_id=(
                        target_image.element_id
                    ),
                    metadata={
                        "relationship": (
                            "embedded_image"
                        )
                    },
                )
            )

    # ---------------------------------------------------------
    # Tables
    # ---------------------------------------------------------
    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):
        rows = []

        for row in table.rows:
            rows.append(
                [
                    cell.text.strip()
                    for cell in row.cells
                ]
            )

        table_text = "\n".join(
            " | ".join(row)
            for row in rows
        )

        element_id = (
            f"table-{table_index}"
        )

        table_element = SourceElement(
            element_id=element_id,
            document_id=document_id,
            element_type="table",
            content=TableContent(
                text=table_text
            ),
            metadata={
                "file_name": docx_path.name,
                "content_type": "table",
                "table_index": table_index,
                "row_count": len(rows),
                "column_count": max(
                    (
                        len(row)
                        for row in rows
                    ),
                    default=0,
                ),
            },
        )

        elements.append(
            table_element
        )

    # ---------------------------------------------------------
    # Root relationship
    # ---------------------------------------------------------
    root_element = SourceElement(
        element_id="document",
        document_id=document_id,
        element_type="document",
        content=TextContent(
            text=""
        ),
        metadata={
            "file_name": docx_path.name,
            "content_type": "document",
        },
    )

    elements.insert(
        0,
        root_element,
    )

    for element in elements[1:]:
        relationships.append(
            ElementRelationship(
                source_element_id="document",
                relationship_type="contains",
                target_element_id=element.element_id,
            )
        )

    return IngestionResult(
        document_id=document_id,
        elements=elements,
        relationships=relationships,
    )