from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import (
    MSO_SHAPE_TYPE,
    PP_PLACEHOLDER,
)

from backend.app.ingestion.models.content import (
    ImageContent,
    TableContent,
    TextContent,
)
from backend.app.ingestion.models.element_relationship import (
    ElementRelationship,
)
from backend.app.ingestion.models.ingestion_result import (
    IngestionResult,
)
from backend.app.ingestion.models.source_element import (
    SourceElement,
)


def _get_bbox(shape) -> dict[str, int]:
    return {
        "left": shape.left,
        "top": shape.top,
        "width": shape.width,
        "height": shape.height,
    }


def _get_shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""

    return shape.text.strip()


def _get_text_element_type(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return "text"

    try:
        placeholder_type = shape.placeholder_format.type

        if placeholder_type in {
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        }:
            return "heading"

        if placeholder_type == PP_PLACEHOLDER.SUBTITLE:
            return "subtitle"

    except ValueError:
        pass

    return "text"


def _get_slide_title(slide) -> str:
    for shape in slide.shapes:
        if not getattr(shape, "is_placeholder", False):
            continue

        try:
            placeholder_type = shape.placeholder_format.type
        except ValueError:
            continue

        if placeholder_type in {
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        }:
            text = _get_shape_text(shape)
            if text:
                return text

    return ""


def _extract_notes(
    slide,
    document_id: str,
    slide_number: int,
) -> SourceElement | None:
    try:
        notes_slide = slide.notes_slide
    except Exception:
        return None

    notes_text_frame = notes_slide.notes_text_frame

    if notes_text_frame is None:
        return None

    text = notes_text_frame.text.strip()

    if not text:
        return None

    return SourceElement(
        element_id=f"slide-{slide_number}-notes",
        document_id=document_id,
        element_type="notes",
        content=TextContent(text=text),
        metadata={
            "slide_number": slide_number,
        },
    )


def load_pptx(
    file_path: str,
    document_id: str,
    output_dir: str,
) -> IngestionResult:
    pptx_path = Path(file_path)

    if not pptx_path.exists():
        raise FileNotFoundError(
            f"PPTX file not found: {file_path}"
        )

    if not pptx_path.is_file():
        raise ValueError(
            f"PPTX path is not a file: {file_path}"
        )

    if pptx_path.suffix.lower() != ".pptx":
        raise ValueError(
            f"Unsupported presentation format: {pptx_path.suffix}"
        )

    output_path = Path(output_dir)
    output_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    presentation = Presentation(str(pptx_path))

    elements: list[SourceElement] = []
    relationships: list[ElementRelationship] = []

    for slide_index, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_element_id = f"slide-{slide_index}"

        slide_element = SourceElement(
            element_id=slide_element_id,
            document_id=document_id,
            element_type="slide",
            content=TextContent(
                text=_get_slide_title(slide)
            ),
            metadata={
                "file_name": pptx_path.name,
                "slide_number": slide_index,
                "content_type": "slide",
            },
        )

        elements.append(slide_element)

        # Extract speaker notes.
        notes_element = _extract_notes(
            slide=slide,
            document_id=document_id,
            slide_number=slide_index,
        )

        if notes_element:
            elements.append(notes_element)

            relationships.append(
                ElementRelationship(
                    source_element_id=slide_element_id,
                    relationship_type="contains",
                    target_element_id=notes_element.element_id,
                    metadata={
                        "slide_number": slide_index,
                    },
                )
            )

        for shape_index, shape in enumerate(
            slide.shapes,
            start=1,
        ):
            base_metadata: dict[str, Any] = {
                "file_name": pptx_path.name,
                "slide_number": slide_index,
                "shape_index": shape_index,
                "bbox": _get_bbox(shape),
            }

            # Text / heading / subtitle
            text = _get_shape_text(shape)

            if text:
                element_type = _get_text_element_type(shape)
                element_id = (
                    f"slide-{slide_index}-"
                    f"{element_type}-{shape_index}"
                )

                text_element = SourceElement(
                    element_id=element_id,
                    document_id=document_id,
                    element_type=element_type,
                    content=TextContent(text=text),
                    metadata={
                        **base_metadata,
                        "content_type": "text",
                    },
                )

                elements.append(text_element)

                relationships.append(
                    ElementRelationship(
                        source_element_id=slide_element_id,
                        relationship_type="contains",
                        target_element_id=element_id,
                        metadata={
                            "slide_number": slide_index,
                        },
                    )
                )

            # Embedded image
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image

                image_extension = image.ext.lower()

                image_path = (
                    output_path
                    / (
                        f"{document_id}_"
                        f"slide_{slide_index}_"
                        f"image_{shape_index}."
                        f"{image_extension}"
                    )
                )

                image_path.write_bytes(image.blob)

                element_id = (
                    f"slide-{slide_index}-"
                    f"image-{shape_index}"
                )

                image_element = SourceElement(
                    element_id=element_id,
                    document_id=document_id,
                    element_type="image",
                    content=ImageContent(
                        path=str(image_path)
                    ),
                    metadata={
                        **base_metadata,
                        "content_type": "image",
                        "image_path": str(image_path),
                        "image_extension": image_extension,
                    },
                )

                elements.append(image_element)

                relationships.append(
                    ElementRelationship(
                        source_element_id=slide_element_id,
                        relationship_type="contains",
                        target_element_id=element_id,
                        metadata={
                            "slide_number": slide_index,
                        },
                    )
                )

            # Table
            if getattr(shape, "has_table", False):
                table = shape.table

                rows = []

                for row in table.rows:
                    rows.append(
                        [
                            cell.text.strip()
                            for cell in row.cells
                        ]
                    )

                table_text = "\n".join(
                    " | ".join(row)
                    for row in rows
                )

                element_id = (
                    f"slide-{slide_index}-"
                    f"table-{shape_index}"
                )

                table_element = SourceElement(
                    element_id=element_id,
                    document_id=document_id,
                    element_type="table",
                    content=TableContent(
                        text=table_text
                    ),
                    metadata={
                        **base_metadata,
                        "content_type": "table",
                        "row_count": len(rows),
                        "column_count": (
                            len(rows[0])
                            if rows
                            else 0
                        ),
                    },
                )

                elements.append(table_element)

                relationships.append(
                    ElementRelationship(
                        source_element_id=slide_element_id,
                        relationship_type="contains",
                        target_element_id=element_id,
                        metadata={
                            "slide_number": slide_index,
                        },
                    )
                )

            # Chart
            if getattr(shape, "has_chart", False):
                chart = shape.chart

                chart_title = ""

                if chart.has_title:
                    chart_title = (
                        chart.chart_title.text_frame.text
                        .strip()
                    )

                element_id = (
                    f"slide-{slide_index}-"
                    f"chart-{shape_index}"
                )

                chart_element = SourceElement(
                    element_id=element_id,
                    document_id=document_id,
                    element_type="chart",
                    content=TextContent(
                        text=chart_title
                    ),
                    metadata={
                        **base_metadata,
                        "content_type": "chart",
                        "chart_title": chart_title,
                    },
                )

                elements.append(chart_element)

                relationships.append(
                    ElementRelationship(
                        source_element_id=slide_element_id,
                        relationship_type="contains",
                        target_element_id=element_id,
                        metadata={
                            "slide_number": slide_index,
                        },
                    )
                )

    return IngestionResult(
        document_id=document_id,
        elements=elements,
        relationships=relationships,
    )